# -*- coding: utf-8 -*-
"""
Road Automation — M19  Report Generator  (Python 3 + openpyxl).

Reads all project output CSVs via project.json and produces a
multi-sheet Excel design report.

Usage:
  python report_generator.py                          uses default project.json
  python report_generator.py <path/to/project.json>  explicit path
"""
from __future__ import annotations

import csv as _csv
import json
import os
import sys
from typing import Any, Dict, List, Optional, Tuple

try:
    from openpyxl import Workbook
    from openpyxl.chart import BarChart, Reference, LineChart
    from openpyxl.chart.series import SeriesLabel
    from openpyxl.styles import (
        Alignment, Border, Font, PatternFill, Side,
    )
    from openpyxl.utils import get_column_letter
except ImportError:
    print(
        "Install openpyxl:  pip install openpyxl",
        file=sys.stderr,
    )
    sys.exit(1)

# ---------------------------------------------------------------------------
# Root / output path convention
# ---------------------------------------------------------------------------

ROOT    = os.path.normpath(os.path.join(os.path.dirname(os.path.abspath(__file__)), ".."))
DEFAULT_PROJECT_JSON = os.path.join(ROOT, "config", "project.json")
DEFAULT_OUT          = os.path.join(ROOT, "out", "report", "design_report.xlsx")

# ---------------------------------------------------------------------------
# Style constants
# ---------------------------------------------------------------------------

_DARK     = "1c2030"
_AMBER    = "F0A500"
_GREEN    = "1c6f44"
_LGREY    = "e8e8e8"
_DGREY    = "555555"
_WHITE    = "FFFFFF"

_H_FILL   = PatternFill("solid", fgColor=_DARK)
_H_FONT   = Font(name="Calibri", size=11, bold=True, color=_AMBER)
_H_ALIGN  = Alignment(horizontal="center", vertical="center", wrap_text=True)

_SUB_FILL = PatternFill("solid", fgColor=_GREEN)
_SUB_FONT = Font(name="Calibri", size=10, bold=True, color=_WHITE)

_B_FONT   = Font(name="Calibri", size=10)
_B_FILL_A = PatternFill("solid", fgColor="f0f4e8")
_B_FILL_B = PatternFill("solid", fgColor=_WHITE)

_THIN     = Side(style="thin", color="cccccc")
_CELL_BDR = Border(left=_THIN, right=_THIN, bottom=_THIN, top=_THIN)

_NUM2     = "0.00"
_NUM1     = "0.0"
_NUM0     = "#,##0"
_PCT1     = "0.0%"


# ---------------------------------------------------------------------------
# CSV reader helper
# ---------------------------------------------------------------------------

def _read_csv(path: str) -> List[Dict[str, str]]:
    if not os.path.isfile(path):
        return []
    with open(path, "r", encoding="utf-8-sig") as f:
        return list(_csv.DictReader(f))


def _flt(row: dict, key: str, default: float = 0.0) -> float:
    try:
        return float(row[key])
    except (KeyError, TypeError, ValueError):
        return default


# ---------------------------------------------------------------------------
# Sheet builders
# ---------------------------------------------------------------------------

def _style_header_row(ws, row: int, ncols: int) -> None:
    for c in range(1, ncols + 1):
        cell = ws.cell(row=row, column=c)
        cell.fill      = _H_FILL
        cell.font      = _H_FONT
        cell.alignment = _H_ALIGN
        cell.border    = _CELL_BDR
    ws.row_dimensions[row].height = 22


def _write_table(
    ws,
    headers:    List[str],
    rows:       List[Dict[str, str]],
    start_row:  int = 2,
    col_offset: int = 1,
) -> int:
    """Write headers + data rows; return the last data row number."""
    ncols = len(headers)
    _style_header_row(ws, start_row, ncols)
    for c_idx, h in enumerate(headers, start=col_offset):
        ws.cell(row=start_row, column=c_idx, value=h)
        ws.column_dimensions[get_column_letter(c_idx)].width = max(len(h) + 4, 12)

    for r_idx, row in enumerate(rows, start=start_row + 1):
        fill = _B_FILL_A if (r_idx % 2 == 0) else _B_FILL_B
        for c_idx, h in enumerate(headers, start=col_offset):
            cell = ws.cell(row=r_idx, column=c_idx, value=row.get(h, ""))
            cell.font   = _B_FONT
            cell.fill   = fill
            cell.border = _CELL_BDR
    return start_row + len(rows)


def _cover_sheet(wb: Workbook, cfg: dict) -> None:
    ws = wb.create_sheet("Cover")
    proj  = cfg.get("project", {})
    names = cfg.get("names",   {})
    ws.column_dimensions["A"].width = 28
    ws.column_dimensions["B"].width = 42

    title = "Road Design Automation — Design Report"
    ws.merge_cells("A1:B1")
    c = ws["A1"]
    c.value     = title
    c.font      = Font(name="Calibri", size=16, bold=True, color=_AMBER)
    c.fill      = _H_FILL
    c.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 36

    data = [
        ("Project Number",  proj.get("number",   "")),
        ("Project Name",    proj.get("name",      "")),
        ("Revision",        proj.get("revision",  "")),
        ("Designer",        proj.get("designer",  "")),
        ("Checker",         proj.get("checker",   "")),
        ("Alignment Name",  names.get("alignment", "")),
        ("Corridor Name",   names.get("corridor",  "")),
    ]
    for row_idx, (label, value) in enumerate(data, start=3):
        ws.cell(row=row_idx, column=1, value=label).font = Font(name="Calibri", size=11, bold=True)
        ws.cell(row=row_idx, column=2, value=value).font = Font(name="Calibri", size=11)
        ws.row_dimensions[row_idx].height = 18

    ws.freeze_panes = None


def _alignment_sheet(wb: Workbook, rows: List[dict]) -> None:
    ws = wb.create_sheet("Alignment")
    if not rows:
        ws["A1"] = "No alignment data (alignment_pi.csv not found)"
        return
    headers = list(rows[0].keys()) if rows else ["pi_id", "easting", "northing", "radius_m"]
    _write_table(ws, headers, rows)
    ws.freeze_panes = "A3"


def _profile_sheet(wb: Workbook, rows: List[dict]) -> None:
    ws = wb.create_sheet("Profile")
    if not rows:
        ws["A1"] = "No profile data (profile_pvis.csv not found)"
        return
    headers = list(rows[0].keys()) if rows else ["station_m", "elevation_m", "k_crest", "k_sag"]
    last_row = _write_table(ws, headers, rows)
    ws.freeze_panes = "A3"

    # Add elevation vs station line chart
    if len(rows) >= 2:
        try:
            chart = LineChart()
            chart.title  = "Vertical Profile"
            chart.y_axis.title = "Elevation (m)"
            chart.x_axis.title = "Station (m)"
            chart.style  = 10
            chart.height = 12
            chart.width  = 20

            sta_col  = headers.index("station_m")  + 1 if "station_m"  in headers else 1
            elev_col = headers.index("elevation_m") + 1 if "elevation_m" in headers else 2

            data_ref  = Reference(ws, min_col=elev_col, min_row=2, max_row=last_row + 1)
            cats_ref  = Reference(ws, min_col=sta_col,  min_row=3, max_row=last_row + 1)
            chart.add_data(data_ref, titles_from_data=True)
            chart.set_categories(cats_ref)
            ws.add_chart(chart, f"A{last_row + 3}")
        except Exception:
            pass


def _sections_sheet(wb: Workbook, rows: List[dict]) -> None:
    ws = wb.create_sheet("Sections")
    if not rows:
        ws["A1"] = "No section data (sections_list.csv not found — run M10)"
        return
    headers = list(rows[0].keys()) if rows else ["station_m", "label", "left_width_m", "right_width_m"]
    _write_table(ws, headers, rows)
    ws.freeze_panes = "A3"


def _volumes_sheet(wb: Workbook, rows: List[dict]) -> None:
    ws = wb.create_sheet("Volumes")
    if not rows:
        ws["A1"] = "No volume data (volumes.csv not found — run M4)"
        return
    headers = list(rows[0].keys())
    _write_table(ws, headers, rows)
    ws.freeze_panes = "A3"

    # Summary row totals
    total_row = len(rows) + 3
    for c_idx, h in enumerate(headers, start=1):
        if h in ("cut_m3", "fill_m3", "net_m3"):
            total = sum(_flt(r, h) for r in rows)
            ws.cell(row=total_row, column=c_idx, value=round(total, 2)).font = Font(bold=True)
    ws.cell(row=total_row, column=1, value="TOTAL").font = Font(bold=True)


def _pavement_sheet(wb: Workbook, rows: List[dict]) -> None:
    ws = wb.create_sheet("Pavement")
    if not rows:
        ws["A1"] = "No pavement data (pavement_design.csv not found — run M16)"
        return
    headers = list(rows[0].keys())
    _write_table(ws, headers, rows)
    ws.freeze_panes = "A3"


def _drainage_sheet(wb: Workbook, rows: List[dict]) -> None:
    ws = wb.create_sheet("Drainage")
    if not rows:
        ws["A1"] = "No drainage data (drainage_design.csv not found — run M17)"
        return
    headers = list(rows[0].keys())
    _write_table(ws, headers, rows)
    ws.freeze_panes = "A3"


def _boq_sheet(wb: Workbook, rows: List[dict]) -> None:
    ws = wb.create_sheet("BOQ")
    if not rows:
        ws["A1"] = "No BOQ data (boq.csv not found — run M7)"
        return
    headers = list(rows[0].keys())
    last_row = _write_table(ws, headers, rows)
    ws.freeze_panes = "A3"

    # Total row for quantity column
    if "quantity" in headers:
        q_col = headers.index("quantity") + 1
        total = sum(_flt(r, "quantity") for r in rows)
        ws.cell(row=last_row + 2, column=q_col, value=round(total, 3)).font = Font(bold=True)
        ws.cell(row=last_row + 2, column=1, value="TOTAL").font = Font(bold=True)


def _mass_haul_sheet(wb: Workbook, rows: List[dict]) -> None:
    ws = wb.create_sheet("MassHaul")
    if not rows:
        ws["A1"] = "No mass haul data (mass_haul.csv not found — run M12)"
        return
    headers = list(rows[0].keys())
    last_row = _write_table(ws, headers, rows)
    ws.freeze_panes = "A3"

    # Mass haul curve chart
    if "station_m" in headers and "cumulative_net_m3" in headers and len(rows) >= 2:
        try:
            chart = LineChart()
            chart.title  = "Mass Haul Diagram"
            chart.y_axis.title = "Cumulative Net Volume (m³)"
            chart.x_axis.title = "Station (m)"
            chart.style  = 10
            chart.height = 14
            chart.width  = 24

            cum_col = headers.index("cumulative_net_m3") + 1
            sta_col = headers.index("station_m") + 1

            data_ref = Reference(ws, min_col=cum_col, min_row=2, max_row=last_row + 1)
            cats_ref = Reference(ws, min_col=sta_col, min_row=3, max_row=last_row + 1)
            chart.add_data(data_ref, titles_from_data=True)
            chart.set_categories(cats_ref)
            ws.add_chart(chart, f"A{last_row + 3}")
        except Exception:
            pass


def _qa_log_sheet(wb: Workbook, qa_log_path: str) -> None:
    ws = wb.create_sheet("QA Log")
    ws.column_dimensions["A"].width = 90
    ws["A1"] = "QA / Preflight Log"
    ws["A1"].font = Font(name="Calibri", size=12, bold=True)

    if not os.path.isfile(qa_log_path):
        ws["A3"] = f"QA log not found: {qa_log_path}"
        ws["A4"] = "Run:  python tools/road_automation_preflight.py"
        return

    with open(qa_log_path, "r", encoding="utf-8") as f:
        lines = f.readlines()

    for row_idx, line in enumerate(lines, start=3):
        cell = ws.cell(row=row_idx, column=1, value=line.rstrip())
        cell.font = Font(name="Consolas", size=9)
        if line.startswith("[ERROR]") or line.startswith("ERROR"):
            cell.font = Font(name="Consolas", size=9, color="CC0000", bold=True)
        elif line.startswith("[WARN]") or line.startswith("WARN"):
            cell.font = Font(name="Consolas", size=9, color="CC6600")
        elif line.startswith("[OK]") or line.startswith("OK"):
            cell.font = Font(name="Consolas", size=9, color=_GREEN)


# ---------------------------------------------------------------------------
# Main
# ---------------------------------------------------------------------------

def generate_report(project_json: str, out_path: str) -> List[str]:
    log: List[str] = []
    log.append(f"Project : {project_json}")
    log.append(f"Output  : {out_path}")

    if not os.path.isfile(project_json):
        log.append(f"[ERROR] project.json not found: {project_json}")
        return log

    with open(project_json, "r", encoding="utf-8-sig") as f:
        cfg = json.load(f)

    base   = os.path.dirname(os.path.abspath(project_json))
    root   = os.path.normpath(os.path.join(base, ".."))
    paths  = cfg.get("paths", {})

    def _p(key: str, default: str) -> str:
        return os.path.normpath(os.path.join(root, paths.get(key, default)))

    align_rows   = _read_csv(_p("alignment_pi",   "csv/alignment_pi.csv"))
    profile_rows = _read_csv(_p("profile_pvis",   "csv/profile_pvis.csv"))
    vol_rows     = _read_csv(_p("volumes_csv",     "out/volumes.csv"))
    boq_rows     = _read_csv(_p("boq_csv",        "out/boq.csv"))
    sect_rows    = _read_csv(os.path.join(root, "out", "sections_list.csv"))
    pave_rows    = _read_csv(os.path.join(root, "out", "pavement_design.csv"))
    drain_rows   = _read_csv(os.path.join(root, "out", "drainage_design.csv"))
    haul_rows    = _read_csv(os.path.join(root, "out", "mass_haul.csv"))
    qa_log_path  = _p("qa_log", "out/qa/run_log.txt")

    wb = Workbook()
    wb.remove(wb.active)

    _cover_sheet(wb, cfg)
    _alignment_sheet(wb, align_rows)
    _profile_sheet(wb, profile_rows)
    _sections_sheet(wb, sect_rows)
    _volumes_sheet(wb, vol_rows)
    _pavement_sheet(wb, pave_rows)
    _drainage_sheet(wb, drain_rows)
    _boq_sheet(wb, boq_rows)
    _mass_haul_sheet(wb, haul_rows)
    _qa_log_sheet(wb, qa_log_path)

    os.makedirs(os.path.dirname(os.path.abspath(out_path)), exist_ok=True)
    wb.save(out_path)

    log.append("")
    log.append("Sheets written:")
    for sh in wb.sheetnames:
        log.append(f"  {sh}")
    log.append(f"\nReport saved: {out_path}")
    return log


# ---------------------------------------------------------------------------
# Entry point
# ---------------------------------------------------------------------------

def main() -> None:
    args = sys.argv[1:]
    project_json = args[0] if args and not args[0].startswith("-") else DEFAULT_PROJECT_JSON
    out_path     = args[1] if len(args) > 1 and not args[1].startswith("-") else DEFAULT_OUT

    for line in generate_report(project_json, out_path):
        print(line)


if __name__ == "__main__":
    main()
